"""
Generative File Creator for AI Girl (Presentations, Spreadsheets, Documents).
Creates downloadable .pptx, .xlsx, and .docx files just like ChatGPT / Copilot.
"""

import os
import re
import uuid
import datetime
from pathlib import Path
from typing import List, Dict, Any, Optional

GENERATED_DIR = Path(__file__).parent.parent / "frontend" / "generated"
GENERATED_DIR.mkdir(parents=True, exist_ok=True)


def sanitize_filename(name: str, fallback: str = "document", ext: str = ".txt") -> str:
    cleaned = re.sub(r'[\\/*?:"<>|]', "", name).strip().replace(" ", "_")
    if not cleaned:
        cleaned = fallback
    if not cleaned.endswith(ext):
        cleaned += ext
    return cleaned


def create_presentation(
    title: str,
    slides_data: List[Dict[str, Any]],
    filename: Optional[str] = None
) -> str:
    """
    Creates a modern 16:9 PowerPoint Presentation (.pptx) file.
    slides_data format:
    [
        {"title": "Introduction", "bullets": ["Point 1", "Point 2", "Point 3"]},
        ...
    ]
    """
    try:
        import pptx
        from pptx import Presentation
        from pptx.util import Inches, Pt
        from pptx.dml.color import RGBColor
        from pptx.enum.text import PP_ALIGN

        prs = Presentation()
        # Set 16:9 Widescreen dimensions
        prs.slide_width = Inches(13.333)
        prs.slide_height = Inches(7.5)

        blank_slide_layout = prs.slide_layouts[6]

        # ── Color Palette (Ice Blue & Modern Tech Theme) ──
        PRIMARY_BLUE = RGBColor(90, 133, 246)    # #5A85F6
        DARK_TEXT = RGBColor(33, 37, 41)         # #212529
        MUTED_TEXT = RGBColor(100, 116, 139)     # #64748B
        ACCENT_PINK = RGBColor(247, 168, 196)    # #F7A8C4
        BG_CARD = RGBColor(243, 248, 255)        # #F3F8FF

        # ── SLIDE 1: Title Slide ──
        title_slide = prs.slides.add_slide(blank_slide_layout)
        
        # Background accent header band
        top_bar = title_slide.shapes.add_shape(
            1, Inches(0), Inches(0), Inches(13.333), Inches(0.4) # 1 = MSO_SHAPE.RECTANGLE
        )
        top_bar.fill.solid()
        top_bar.fill.fore_color.rgb = PRIMARY_BLUE
        top_bar.line.fill.background()

        # Title Text Box
        t_box = title_slide.shapes.add_textbox(Inches(1.5), Inches(2.2), Inches(10.33), Inches(3.0))
        tf = t_box.text_frame
        tf.word_wrap = True

        p_main = tf.paragraphs[0]
        p_main.text = title
        p_main.font.size = Pt(44)
        p_main.font.bold = True
        p_main.font.color.rgb = PRIMARY_BLUE
        p_main.font.name = "Arial"

        p_sub = tf.add_paragraph()
        p_sub.text = f"Prepared by Prity AI Companion • {datetime.datetime.now().strftime('%B %Y')}"
        p_sub.font.size = Pt(20)
        p_sub.font.color.rgb = MUTED_TEXT
        p_sub.font.name = "Arial"
        p_sub.space_before = Pt(20)

        # ── CONTENT SLIDES ──
        for s_idx, slide_info in enumerate(slides_data, start=2):
            slide = prs.slides.add_slide(blank_slide_layout)

            # Slide Top Bar
            s_top_bar = slide.shapes.add_shape(
                1, Inches(0), Inches(0), Inches(13.333), Inches(0.2)
            )
            s_top_bar.fill.solid()
            s_top_bar.fill.fore_color.rgb = PRIMARY_BLUE
            s_top_bar.line.fill.background()

            # Slide Title
            header_box = slide.shapes.add_textbox(Inches(1.0), Inches(0.7), Inches(11.33), Inches(1.0))
            h_tf = header_box.text_frame
            h_tf.word_wrap = True
            h_p = h_tf.paragraphs[0]
            h_p.text = slide_info.get("title", f"Slide {s_idx}")
            h_p.font.size = Pt(30)
            h_p.font.bold = True
            h_p.font.color.rgb = PRIMARY_BLUE
            h_p.font.name = "Arial"

            # Content Card
            content_card = slide.shapes.add_shape(
                1, Inches(1.0), Inches(1.9), Inches(11.33), Inches(4.8)
            )
            content_card.fill.solid()
            content_card.fill.fore_color.rgb = BG_CARD
            content_card.line.color.rgb = RGBColor(220, 233, 248)

            # Content Box (Bullets or text)
            c_box = slide.shapes.add_textbox(Inches(1.4), Inches(2.2), Inches(10.5), Inches(4.2))
            c_tf = c_box.text_frame
            c_tf.word_wrap = True

            bullets = slide_info.get("bullets", [])
            content_text = slide_info.get("content", "")

            if bullets:
                for b_idx, bullet in enumerate(bullets):
                    bp = c_tf.paragraphs[0] if b_idx == 0 else c_tf.add_paragraph()
                    bp.text = f"•  {bullet}"
                    bp.font.size = Pt(19)
                    bp.font.color.rgb = DARK_TEXT
                    bp.font.name = "Arial"
                    bp.space_after = Pt(14)
            elif content_text:
                cp = c_tf.paragraphs[0]
                cp.text = content_text
                cp.font.size = Pt(20)
                cp.font.color.rgb = DARK_TEXT
                cp.font.name = "Arial"

            # Slide Number Footer
            f_box = slide.shapes.add_textbox(Inches(11.0), Inches(6.8), Inches(1.5), Inches(0.5))
            f_tf = f_box.text_frame
            f_p = f_tf.paragraphs[0]
            f_p.text = f"{s_idx - 1} / {len(slides_data)}"
            f_p.font.size = Pt(12)
            f_p.font.color.rgb = MUTED_TEXT

        if not filename:
            safe_title = sanitize_filename(title[:30], fallback="Presentation", ext=".pptx")
            filename = f"{safe_title}"

        final_path = GENERATED_DIR / filename
        prs.save(str(final_path))
        return filename
    except Exception as e:
        print(f"Error creating PPTX: {e}")
        return ""


def create_excel_workbook(
    sheet_title: str,
    headers: List[str],
    rows_data: List[List[Any]],
    filename: Optional[str] = None
) -> str:
    """Creates a styled Excel Spreadsheet (.xlsx) file."""
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = (sheet_title or "Data")[:31]

        # Header Style
        header_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="5A85F6", end_color="5A85F6", fill_type="solid")
        header_align = Alignment(horizontal="center", vertical="center")

        # Zebra Rows
        row_fill_alt = PatternFill(start_color="F3F8FF", end_color="F3F8FF", fill_type="solid")
        thin_border = Border(
            left=Side(style="thin", color="E2E8F0"),
            right=Side(style="thin", color="E2E8F0"),
            top=Side(style="thin", color="E2E8F0"),
            bottom=Side(style="thin", color="E2E8F0")
        )

        # Write Headers
        if headers:
            ws.append(headers)
            for col_num in range(1, len(headers) + 1):
                cell = ws.cell(row=1, column=col_num)
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = header_align
                cell.border = thin_border

        # Write Data Rows
        for r_idx, row in enumerate(rows_data, start=2 if headers else 1):
            ws.append(row)
            for col_num in range(1, len(row) + 1):
                cell = ws.cell(row=r_idx, column=col_num)
                cell.font = Font(name="Arial", size=10)
                cell.border = thin_border
                if r_idx % 2 == 1:
                    cell.fill = row_fill_alt

        # Auto-adjust column widths
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val_str = str(cell.value or "")
                if len(val_str) > max_len:
                    max_len = len(val_str)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

        if not filename:
            safe_title = sanitize_filename(sheet_title[:30], fallback="Spreadsheet", ext=".xlsx")
            filename = safe_title

        final_path = GENERATED_DIR / filename
        wb.save(str(final_path))
        return filename
    except Exception as e:
        print(f"Error creating Excel: {e}")
        return ""


def create_word_doc(
    title: str,
    sections: List[Dict[str, Any]],
    filename: Optional[str] = None
) -> str:
    """Creates a styled Word Document (.docx) file."""
    try:
        import docx
        from docx.shared import Inches, Pt, RGBColor

        doc = docx.Document()

        # Document Title
        title_p = doc.add_paragraph()
        title_run = title_p.add_run(title)
        title_run.font.size = Pt(26)
        title_run.font.bold = True
        title_run.font.color.rgb = RGBColor(90, 133, 246)

        doc.add_paragraph(f"Generated by Prity AI Companion • {datetime.datetime.now().strftime('%B %d, %Y')}")
        doc.add_paragraph("-" * 60)

        for sec in sections:
            sec_title = sec.get("heading", "")
            if sec_title:
                h = doc.add_heading(sec_title, level=2)
                for run in h.runs:
                    run.font.color.rgb = RGBColor(45, 55, 72)

            body = sec.get("text", "")
            if body:
                doc.add_paragraph(body)

            bullets = sec.get("bullets", [])
            for bullet in bullets:
                doc.add_paragraph(bullet, style="List Bullet")

            table_data = sec.get("table", [])
            if table_data:
                table = doc.add_table(rows=len(table_data), cols=len(table_data[0]))
                table.style = 'Table Grid'
                for r_idx, row in enumerate(table_data):
                    for c_idx, val in enumerate(row):
                        table.cell(r_idx, c_idx).text = str(val)

        if not filename:
            safe_title = sanitize_filename(title[:30], fallback="Document", ext=".docx")
            filename = safe_title

        final_path = GENERATED_DIR / filename
        doc.save(str(final_path))
        return filename
    except Exception as e:
        print(f"Error creating Word doc: {e}")
        return ""


def auto_detect_and_generate_files(user_msg: str, ai_reply: str) -> str:
    """
    Scans the conversation to see if the user requested a presentation (PPT),
    spreadsheet (Excel), or document (Word). Automatically creates the file
    and attaches an interactive download widget.
    """
    lower_user = user_msg.lower()
    lower_reply = ai_reply.lower()

    # 1. PowerPoint Presentation Generation
    ppt_triggers = ["presentation", "powerpoint", "ppt", "slides", "pptx", "slide deck", "make a ppt", "create a ppt"]
    if any(trig in lower_user for trig in ppt_triggers) and ("slide 1" in lower_reply or "slide" in lower_reply or "#" in ai_reply):
        try:
            # Extract slide titles and bullet points from reply
            slides = []
            current_slide = None

            lines = ai_reply.splitlines()
            for line in lines:
                line_str = line.strip()
                if not line_str:
                    continue
                # Match Slide header e.g. "Slide 1: Title", "## Slide 1 - Title", "### Title"
                slide_match = re.match(r'^(?:#+\s*|)(?:slide\s*\d+[\s:\-–—]+|)(.+)$', line_str, re.IGNORECASE)
                if ("slide" in line_str.lower() and (":" in line_str or "-" in line_str)) or line_str.startswith("#"):
                    title_candidate = re.sub(r'^[#\s\-*]+', '', line_str).replace('**', '').strip()
                    if current_slide:
                        slides.append(current_slide)
                    current_slide = {"title": title_candidate, "bullets": []}
                elif line_str.startswith(("-", "*", "•", "1.", "2.", "3.", "4.", "5.")) and current_slide:
                    cleaned_bullet = re.sub(r'^[\-\*•\d\.\s]+', '', line_str).replace('**', '').strip()
                    if cleaned_bullet:
                        current_slide["bullets"].append(cleaned_bullet)
                elif current_slide and not current_slide.get("bullets"):
                    cleaned = line_str.replace('**', '').strip()
                    if cleaned:
                        current_slide["bullets"].append(cleaned)

            if current_slide:
                slides.append(current_slide)

            if len(slides) >= 2:
                # Generate PPT
                doc_title = slides[0].get("title", "AI Presentation")
                fname = create_presentation(doc_title, slides)
                if fname:
                    download_card = (
                        f"\n\n[FILE_CARD:pptx:/api/download/file/{fname}:{doc_title} ({len(slides)} Slides):PowerPoint Presentation Ready to Download]"
                    )
                    return ai_reply + download_card
        except Exception as err:
            print(f"Auto-generate PPT error: {err}")

    # 2. Excel Spreadsheet Generation
    excel_triggers = ["excel", "spreadsheet", "xlsx", "csv", "table of", "budget sheet", "create an excel", "make an excel"]
    if any(trig in lower_user for trig in excel_triggers) and ("|" in ai_reply):
        try:
            # Parse Markdown tables from reply
            table_lines = [l.strip() for l in ai_reply.splitlines() if l.strip().startswith("|") and l.strip().endswith("|")]
            if len(table_lines) >= 3:
                header_raw = table_lines[0].strip("|").split("|")
                headers = [h.strip() for h in header_raw]
                data_rows = []
                for row_line in table_lines[2:]: # skip separator
                    cells = [c.strip() for c in row_line.strip("|").split("|")]
                    if any(cells):
                        data_rows.append(cells)

                if headers and data_rows:
                    sheet_name = "AI Data Sheet"
                    fname = create_excel_workbook(sheet_name, headers, data_rows)
                    if fname:
                        download_card = (
                            f"\n\n[FILE_CARD:xlsx:/api/download/file/{fname}:{sheet_name} ({len(data_rows)} Rows):Excel Spreadsheet File Ready to Download]"
                        )
                        return ai_reply + download_card
        except Exception as err:
            print(f"Auto-generate Excel error: {err}")

    # 3. Word Document Generation
    doc_triggers = ["word document", "docx", "write a document", "create a doc", "generate a docx", "make a word doc"]
    if any(trig in lower_user for trig in doc_triggers):
        try:
            sections = []
            curr_sec = {"heading": "Overview", "text": "", "bullets": []}
            for line in ai_reply.splitlines():
                l = line.strip()
                if l.startswith("#"):
                    if curr_sec["text"] or curr_sec["bullets"]:
                        sections.append(curr_sec)
                    curr_sec = {"heading": l.lstrip("#").strip(), "text": "", "bullets": []}
                elif l.startswith(("-", "*", "•")):
                    curr_sec["bullets"].append(l.lstrip("-*• ").strip())
                elif l:
                    curr_sec["text"] += (l + "\n")
            if curr_sec["text"] or curr_sec["bullets"]:
                sections.append(curr_sec)

            if sections:
                doc_title = "Document"
                fname = create_word_doc(doc_title, sections)
                if fname:
                    download_card = (
                        f"\n\n[FILE_CARD:docx:/api/download/file/{fname}:{doc_title}:Word Document Ready to Download]"
                    )
                    return ai_reply + download_card
        except Exception as err:
            print(f"Auto-generate Word error: {err}")

    return ai_reply
